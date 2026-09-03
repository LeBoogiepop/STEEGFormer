# -*- coding: utf-8 -*-
"""Assemble le rapport Markdown en .docx A4 (pandoc + python-docx)."""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "Lacombe_Maxime_Rapport_Stage_ESME_INGE3.docx"
MASTER = ROOT / "_MASTER_BUILD.md"

PAGEBREAK = """
```{=openxml}
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
```

"""

ORDER = [
    "00_page_de_garde.md",
    "01_introduction.md",
    "02_contexte_labo.md",
    "03_etat_de_lart.md",
    "04_methodes.md",
    "05_travail_realise.md",
    "06_resultats.md",
    "07_discussion.md",
    "08_conclusion.md",
    "BIBLIO.md",
    "ANNEXES.md",
]


TOC_BLOCK = r'''
**Table des matières**

```{=openxml}
<w:p>
  <w:r><w:fldChar w:fldCharType="begin"/></w:r>
  <w:r><w:instrText xml:space="preserve"> TOC \o "1-1" \h \z \u </w:instrText></w:r>
  <w:r><w:fldChar w:fldCharType="separate"/></w:r>
  <w:r><w:t>Clic droit ici, Mettre a jour le champ, puis Mettre a jour toute la table.</w:t></w:r>
  <w:r><w:fldChar w:fldCharType="end"/></w:r>
</w:p>
```

''' + PAGEBREAK


def clean_cover(text: str) -> str:
    text = text.replace("# Page de garde (brouillon Markdown)\n\n", "")
    text = re.sub(
        r"^# Évaluation de ST-EEGFormer[^\n]*",
        "**Évaluation de ST-EEGFormer sur une tâche d'attention spatiale EEG**",
        text,
        count=1,
        flags=re.M,
    )
    text = re.sub(
        r"^### Rapport de stage de fin d.études",
        "*Rapport de stage de fin d’études*",
        text,
        count=1,
        flags=re.M,
    )
    text = re.sub(
        r"\n\*Logos ESME[^\n]*\n",
        "\n*Brouillon à viser par le laboratoire avant dépôt Moodle (14 septembre 2026).*\n",
        text,
        count=1,
    )
    text = re.sub(
        r"\n---\n\n## Table des matières \(prévisionnelle\).*",
        "",
        text,
        flags=re.S,
    )
    return text.strip() + "\n\n" + PAGEBREAK + TOC_BLOCK


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def postprocess(path: Path) -> None:
    doc = Document(str(path))
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = "Maxime Lacombe — Rapport de stage ESME INGE3 — Ishii Lab, Kyoto University"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in hp.runs:
            run.font.size = Pt(8)
            run.font.name = "Calibri"
            run.font.color.rgb = None
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = ""
        r = fp.add_run("Page ")
        r.font.size = Pt(9)
        add_page_number(fp)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    # Demander à Word de rafraîchir les champs (TOC) à l'ouverture.
    settings = doc.settings.element
    already = settings.find(qn("w:updateFields"))
    if already is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
    for para in doc.paragraphs[:8]:
        if "Évaluation de ST-EEGFormer" in para.text:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.size = Pt(18)
                run.bold = True
            break
    doc.save(str(path))


def update_toc_with_word(path: Path) -> None:
    """Remplit le champ TOC (besoin de Word installé)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 absent : ouvrir le .docx et mettre a jour la table des matieres.")
        return
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None
    try:
        doc = word.Documents.Open(str(path), ReadOnly=False)
        doc.Fields.Update()
        if doc.TablesOfContents.Count >= 1:
            doc.TablesOfContents(1).Update()
        doc.Save()
        print("TOC mise a jour via Word.")
    except Exception as exc:
        print("TOC non mise a jour automatiquement:", exc)
        print("Ouvrir le .docx, clic droit sur la table -> Mettre a jour le champ.")
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        word.Quit()


def main() -> None:
    parts: list[str] = [
        "---\nlang: fr\n---\n\n",
    ]
    for name in ORDER:
        p = ROOT / name
        raw = p.read_text(encoding="utf-8")
        if name == "00_page_de_garde.md":
            parts.append(clean_cover(raw))
        else:
            parts.append(raw.rstrip() + "\n\n" + PAGEBREAK)
    MASTER.write_text("".join(parts), encoding="utf-8")

    cmd = [
        "pandoc",
        str(MASTER),
        "-o",
        str(OUT_DOCX),
        "--from",
        "markdown+raw_attribute",
        "--to",
        "docx",
        "--resource-path",
        str(ROOT),
        "--metadata",
        "lang=fr",
    ]
    subprocess.check_call(cmd, cwd=str(ROOT))
    postprocess(OUT_DOCX)
    update_toc_with_word(OUT_DOCX)
    print("Wrote", OUT_DOCX, "size", OUT_DOCX.stat().st_size)


if __name__ == "__main__":
    main()
